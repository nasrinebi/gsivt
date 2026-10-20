"""Read computed results from a recalculated workbook (cached values), for comparison with the engine."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from .model import LEVELS, WQ_BASES, _header_row
from .engine import SENS_RATES, OM_MULTS

CALC_HEADERS = {
    "Row check": "check",
    "Runoff coefficient": "rv",
    "Runoff depth (in/yr)": "rd",
    "Annual runoff (gal/yr)": "rg",
    "Share of runoff captured": "cap",
    "Crediting applied": "dapp",
    "Infiltrating (1/0)": "infd",
    "Non-DCIA CN": "cn",
    "Retention table efficiency": "reff",
    "Volume treated (gal/yr)": "vt",
    "Volume retained (gal/yr)": "vr",
    "TN load reduction (%)": "etn",
    "TP load reduction (%)": "etp",
    "TSS load reduction (%)": "etss",
    "TN load (lb/yr)": "tnl",
    "TN removed (lb/yr)": "tnr",
    "TP load (lb/yr)": "tpl",
    "TP removed (lb/yr)": "tpr",
    "TSS load (lb/yr)": "tssl",
    "TSS removed (lb/yr)": "tssr",
    "Stormwater $/yr": "sw",
    "Water quality $/yr": "wq",
    "CO2 uptake (t/yr)": "co2",
    "Carbon $/yr": "carb",
    "Air quality $/yr": "air",
    "Energy $/yr": "en",
    "Recreation visits (yr 1)": "vis",
    "Recreation $/yr (yr 1)": "rec",
    "Annual benefit $ (yr 1)": "ben",
    "Cost basis": "basis",
    "Cost quantity (sq ft or imp. ac)": "qty",
    "Scale factor": "sf",
    "Construction cost $": "cons",
    "Upfront capital $ (0 if existing)": "upf",
    "Annual O&M $": "om",
    "Service life (yr)": "life",
    "First replacement (yr)": "t1",
    "Replacement PV $": "rep",
    "Cost NPV $": "cnpv",
    "Benefit NPV $": "bnpv",
    "Net NPV $": "net",
    "BCR": "bcr",
    "Annualized cost $/yr": "ann",
    "Annualized cost per lb TN": "ctn",
    "Annualized cost per lb TP": "ctp",
    "Construction $ per lb TN/yr": "captn",
    "Habitat created (ac)": "hab",
    "WQ storage volume (gal)": "wqv",
    "Equivalent grey storage capital $": "grey",
    "In strategy A (1/0)": "ina",
    "In strategy B (1/0)": "inb",
    "Included #": "idx",
    "Vegetated area (sq ft)": "vega",
    "Trees (count)": "trees",
    "Herbaceous air area (sq ft)": "herb",
    "Green roof energy area (sq ft)": "gra",
    "Chart label": "lbl",
    "Crediting method": "meth",
    "CN lower row": "cnlo",
    "CN weight": "wcn",
    "DCIA lower col (App. N)": "dlo",
    "DCIA weight (App. N)": "wd",
    "DCIA lower col (App. O)": "dlo2",
    "DCIA weight (App. O)": "wd2",
    "Residence time used (d)": "tday",
    "TN value $/yr": "tnv",
    "TP value $/yr": "tpv",
    "TSS value $/yr": "tssv",
    "Largest pollutant value $/yr": "wqmax",
    "WQ $/yr, Conservative": "wqc",
    "WQ $/yr, Typical": "wqt",
    "WQ $/yr, High": "wqh",
    "WQ $/yr, other dollar-value source": "wqalt",
}


def read_results(path: str | Path) -> dict:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb["Calculator"]
    hr = _header_row(ws)
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(hr, c).value
        if h in CALC_HEADERS:
            cols[CALC_HEADERS[h]] = c
    rows = [{k: ws.cell(r, c).value for k, c in cols.items()} for r in range(hr + 1, hr + 13)]
    total = {k: ws.cell(hr + 13, c).value for k, c in cols.items()}
    total["area"] = ws.cell(hr + 13, 7).value
    total["fp"] = ws.cell(hr + 13, 9).value
    sc = wb["Scenarios"]
    v = lambda ref: sc[ref].value
    lv = {}
    for j, lev in enumerate(LEVELS):
        cl = "BCD"[j]
        lv[lev] = dict(parts=[v(f"{cl}{r}") for r in range(6, 12)], ben=v(f"{cl}12"), bnpv=v(f"{cl}13"), cnpv=v(f"{cl}14"),
                       net=v(f"{cl}15"), bcr=v(f"{cl}16"), jobs=v(f"{cl}17"))
    dr = {rate: dict(bnpv=v(f"B{22+i}"), cnpv=v(f"C{22+i}"), net=v(f"D{22+i}"), bcr=v(f"E{22+i}")) for i, rate in enumerate(SENS_RATES)}
    st = {}
    for j, lab in enumerate(("A", "B", "All")):
        cl = "BCD"[j]
        keys = ("n", "area", "tnr", "tpr", "vr", "upf", "cnpv", "bnpv")
        d = {k: v(f"{cl}{31+i}") for i, k in enumerate(keys)}
        d.update(net=v(f"{cl}39"), bcr=v(f"{cl}40"), ctn=v(f"{cl}41"))
        st[lab] = d
    wqb = {b: dict(wq=v(f"B{46+i}"), bnpv=v(f"C{46+i}"), cnpv=v(f"D{46+i}"), bcr=v(f"E{46+i}")) for i, b in enumerate(WQ_BASES)}
    om = {m: dict(om=v(f"B{54+i}"), bnpv=v(f"C{54+i}"), cnpv=v(f"D{54+i}"), bcr=v(f"E{54+i}")) for i, m in enumerate(OM_MULTS)}
    vs = dict(selected=dict(wq=v("B61"), bnpv=v("C61"), bcr=v("E61")), other=dict(wq=v("B62"), bnpv=v("C62"), bcr=v("E62")))
    return dict(rows=rows, total=total, scenarios=dict(levels=lv, discount=dr, strategies=st, wq_basis=wqb, om=om, value_source=vs))


def compare(engine: dict, wb: dict, rel_tol: float = 1e-9, abs_tol: float = 1e-6) -> dict:
    """Compare every numeric and text quantity present in both result sets.
    Returns the number of compared values, the largest relative difference and the list of mismatches."""
    out = dict(n=0, max_rel=0.0, mismatches=[])

    def cmp(a, b, where):
        if isinstance(b, bool) or b is None:
            return
        if isinstance(a, str) or isinstance(b, str):
            out["n"] += 1
            if str(a) != str(b):
                out["mismatches"].append((where, a, b))
            return
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            out["n"] += 1
            d = abs(a - b)
            rel = d / max(abs(a), abs(b)) if max(abs(a), abs(b)) > 0 else 0.0
            if max(abs(a), abs(b)) > 1e-9 and rel > out["max_rel"]:
                out["max_rel"] = rel
                out["max_rel_at"] = where
            if d > abs_tol and rel > rel_tol:
                out["mismatches"].append((where, a, b))

    def walk(a, b, where):
        if isinstance(b, dict):
            for k in b:
                if isinstance(a, dict) and k in a:
                    walk(a[k], b[k], f"{where}.{k}")
        elif isinstance(b, list) and isinstance(a, list):
            for i, (x, y) in enumerate(zip(a, b)):
                walk(x, y, f"{where}[{i}]")
        else:
            cmp(a, b, where)

    walk(engine["rows"], wb["rows"], "rows")
    walk(engine["total"], wb["total"], "total")
    walk(engine["scenarios"], wb["scenarios"], "scenarios")
    return out
