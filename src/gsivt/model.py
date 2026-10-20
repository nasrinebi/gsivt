"""Read the GSI Monetary Valuation Tool workbook (v3.3) into Python structures.

The workbook is the single source of data: rule tables (AH Vol. I Appendices N and O), land-use EMCs,
rainfall stations, county zones, soils, per-BMP coefficients, valuation coefficients and the CPI-U table
are all read from it, so edits a user makes on the reference sheets are honoured.

Base coefficient values (columns B to D and their dollar years R to T of the Coefficients sheet) are read as
stored values. Some of them are constant formulas such as =ROUND(814*CRF,2); reading those requires a file
that has been saved by a spreadsheet application (cached values present). Everything that depends on the
Calculator settings (CPI adjustment to the base year, the valuation level, the dollar-value source,
the per-BMP CPI factors) is recomputed here and is never taken from the workbook.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

RET = "Retention (AH App. O tables)"
WD = "Wet detention (AH App. O equation)"
FIX = "Fixed efficiency"
INF = "Infiltrating (no underdrain)"
FIL = "Filtering (underdrain)"
BAM = "Filtering with BAM media"
IMP_BASIS = "Impervious acre treated"
FL_SOURCE = "Florida BMAP projects (2025)"
FL_COSTS = "Florida BMAP medians (ponds, wetland)"
SIMPLE = "Simple Method (Schueler 1987)"
LEVELS = ("Conservative", "Typical", "High")
WQ_BASES = ("Largest single pollutant", "Sum of TN, TP and TSS", "TN only", "TP only")

SETTING_NAMES = ("Station", "County", "RunoffMethod", "Level", "WQBasis", "ValueSource", "IncludeRunoff", "BaseYear",
                 "Period", "DiscRate", "VisitGrowth", "IncludeRec", "DesignDepth", "PermitPct", "LandPerAcre", "EOS",
                 "CostIndex", "CostBasis")
ROW_FIELDS = ("inc", "strat", "bmp", "status", "age", "lu", "area", "imp", "fp", "design", "rdep", "rtime", "soil", "access")
BMP_FIELDS = ("name", "short", "method", "always", "tn", "tp", "tss", "basis", "capital", "om", "life", "veg", "tree",
              "energy", "habitat", "note", "capyr", "omyr", "capfac", "omfac", "bam", "flcap", "flom")


def _ref(wb, name: str):
    dn = wb.defined_names[name]
    sheet, rng = next(iter(dn.destinations))
    return wb[sheet], rng.replace("$", "")


def _cells(wb, name: str):
    ws, rng = _ref(wb, name)
    cells = ws[rng]
    if not isinstance(cells, tuple):
        return [[cells]]
    if not isinstance(cells[0], tuple):
        return [list(cells)]
    return [list(r) for r in cells]


@dataclass
class Model:
    """Reference data and coefficients of one workbook."""
    roc: dict = field(default_factory=dict)          # (zone, cn) -> {dcia: value}
    ret: dict = field(default_factory=dict)          # (zone, 'd.dd', cn) -> {dcia: value %}
    landuse: dict = field(default_factory=dict)      # name -> (tn, tp, tss)
    stations: dict = field(default_factory=dict)     # name -> rainfall in/yr
    counties: dict = field(default_factory=dict)     # name -> zone
    soils: dict = field(default_factory=dict)        # label -> CN
    bmps: dict = field(default_factory=dict)         # name -> dict of BMP_FIELDS
    coef: dict = field(default_factory=dict)         # name -> dict(cons, typ, high, ycons, ytyp, yhigh)
    cpi: dict = field(default_factory=dict)          # year -> CPI-U
    lists: dict = field(default_factory=dict)        # list name -> values
    source: str = ""

    @classmethod
    def from_workbook(cls, path: str | Path) -> "Model":
        path = str(path)
        wbv = openpyxl.load_workbook(path, data_only=True)
        m = cls(source=path)
        # Appendix N runoff coefficients
        hdr = [c.value for c in _cells(wbv, "ROCHdr")[0]]
        for key, row in zip(_cells(wbv, "ROCKeys"), _cells(wbv, "ROCData")):
            z, cn = key[0].value.split("|")
            m.roc[(int(z), int(cn))] = {int(h): c.value for h, c in zip(hdr, row)}
        # Appendix O retention tables
        hdr = [c.value for c in _cells(wbv, "RetHdr")[0]]
        for key, row in zip(_cells(wbv, "RetKeys"), _cells(wbv, "RetData")):
            z, d, cn = key[0].value.split("|")
            m.ret[(int(z), d, int(cn))] = {int(h): c.value for h, c in zip(hdr, row)}
        m.landuse = {r[0].value: (r[1].value, r[2].value, r[3].value) for r in _cells(wbv, "LandUseTable")}
        m.stations = {r[0].value: r[1].value for r in _cells(wbv, "StationTable")}
        m.counties = {r[0].value: r[1].value for r in _cells(wbv, "CountyTable")}
        m.soils = {r[0].value: r[1].value for r in _cells(wbv, "SoilTable")}
        m.cpi = {r[0].value: r[1].value for r in _cells(wbv, "CPITable") if isinstance(r[0].value, (int, float))}
        for r in _cells(wbv, "BMPTable"):
            d = {k: c.value for k, c in zip(BMP_FIELDS, r)}
            m.bmps[d["name"]] = d
        # valuation coefficients: every named cell in column E of the Coefficients sheet
        ws = wbv["Coefficients"]
        for nm in wbv.defined_names:
            sh, rng = _ref(wbv, nm)
            if sh.title == "Coefficients" and rng.startswith("E") and ":" not in rng:
                r = int(rng[1:])
                m.coef[nm] = dict(cons=ws.cell(r, 2).value, typ=ws.cell(r, 3).value, high=ws.cell(r, 4).value,
                                  ycons=ws.cell(r, 18).value, ytyp=ws.cell(r, 19).value, yhigh=ws.cell(r, 20).value)
        for nm in ("YesNo", "StrategyList", "StatusList", "DesignList", "LevelList", "PeriodList", "RateList", "DepthList",
                   "GrowthList", "OnOff", "PermitList", "SensRates", "RetDepthList", "MethodList", "BaseYearList",
                   "WQBasisList", "ValueSourceList", "CostBasisList"):
            m.lists[nm] = [c.value for row in _cells(wbv, nm) for c in row]
        return m

    # ------------------------------------------------------------------ CPI helpers (IFERROR semantics of the workbook)
    def cpi_factor(self, base_year: Any, year: Any) -> float:
        try:
            return self.cpi[base_year] / self.cpi[year]
        except (KeyError, TypeError, ZeroDivisionError):
            return 1.0


def read_inputs(path: str | Path) -> tuple[dict, list[dict]]:
    """Settings and the 12 Calculator rows as entered in a workbook (input cells only)."""
    wb = openpyxl.load_workbook(str(path), data_only=False)
    s = {}
    for nm in SETTING_NAMES:
        ws, rng = _ref(wb, nm)
        s[nm] = ws[rng].value
    s["ProjectName"] = _cells(wb, "ProjectName")[0][0].value
    ws = wb["Calculator"]
    hr = _header_row(ws)
    rows = []
    for r in range(hr + 1, hr + 13):
        rows.append({k: ws.cell(r, j + 1).value for j, k in enumerate(ROW_FIELDS)})
    return s, rows


def _header_row(ws) -> int:
    for r in range(1, 60):
        if ws.cell(r, 1).value == "Include?":
            return r
    raise ValueError("Calculator header row not found")


def write_inputs(src: str | Path, dst: str | Path, settings: dict, rows: list[dict]) -> None:
    """Copy a workbook and replace its Calculator inputs (formulas are kept; recalculate before reading results)."""
    wb = openpyxl.load_workbook(str(src))
    for nm, v in settings.items():
        if nm in SETTING_NAMES or nm == "ProjectName":
            ws, rng = _ref(wb, nm)
            ws[rng].value = v
    ws = wb["Calculator"]
    hr = _header_row(ws)
    for i, row in enumerate(rows):
        for j, k in enumerate(ROW_FIELDS):
            ws.cell(hr + 1 + i, j + 1).value = row[k]
    wb.save(str(dst))
